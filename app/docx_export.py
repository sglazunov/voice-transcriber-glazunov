"""Generate a Word .docx report from a transcription job."""
from __future__ import annotations

import datetime
from pathlib import Path
from typing import List, Optional


def _fmt_time(seconds: float) -> str:
    s = int(seconds)
    h, rem = divmod(s, 3600)
    m, sec = divmod(rem, 60)
    if h:
        return f"{h:02d}:{m:02d}:{sec:02d}"
    return f"{m:02d}:{sec:02d}"


def generate_report(
    out_path: Path,
    filename: str,
    segments: list,
    analysis: dict,
    duration: Optional[float] = None,
    date_str: Optional[str] = None,
) -> None:
    """Write a Word document with summary, key thoughts, tasks and full transcript."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        raise RuntimeError(
            "Пакет python-docx не установлен. Выполните: pip install python-docx"
        )

    doc = Document()

    # ---- page margins (narrower for readability) ----------------------------
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # ---- default paragraph font ---------------------------------------------
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    # ---- title block --------------------------------------------------------
    title_p = doc.add_heading("Протокол встречи", level=0)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(Path(filename).stem)
    run.bold = True
    run.font.size = Pt(12)

    meta_line = doc.add_paragraph()
    meta_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shown_date = date_str or datetime.datetime.now().strftime("%d.%m.%Y")
    dur_str = f"  ·  Длительность: {_fmt_time(duration)}" if duration else ""
    meta_run = meta_line.add_run(f"Дата: {shown_date}{dur_str}")
    meta_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    meta_run.font.size = Pt(10)

    doc.add_paragraph()  # spacer

    # ---- О чём шёл разговор -------------------------------------------------
    doc.add_heading("О чём шёл разговор", level=1)
    p = doc.add_paragraph(analysis.get("summary", ""))
    p.paragraph_format.space_after = Pt(12)

    # ---- Подробный разбор по темам -----------------------------------------
    detailed = analysis.get("detailed", [])
    if detailed:
        doc.add_heading("Подробный разбор", level=1)
        for block in detailed:
            topic = (block.get("topic") or "").strip()
            details = (block.get("details") or "").strip()
            if topic:
                doc.add_heading(topic, level=2)
            if details:
                dp = doc.add_paragraph(details)
                dp.paragraph_format.space_after = Pt(10)

    # ---- Ключевые мысли -----------------------------------------------------
    doc.add_heading("Ключевые мысли", level=1)
    thoughts = analysis.get("key_thoughts", [])
    if thoughts:
        for thought in thoughts:
            bp = doc.add_paragraph(style="List Bullet")
            bp.add_run(thought)
    else:
        doc.add_paragraph("—").paragraph_format.space_after = Pt(6)

    # ---- Выводы -------------------------------------------------------------
    conclusions = analysis.get("conclusions", [])
    if conclusions:
        doc.add_heading("Выводы", level=1)
        for c in conclusions:
            bp = doc.add_paragraph(style="List Bullet")
            bp.add_run(c)

    # ---- Решения / договорённости ------------------------------------------
    decisions = analysis.get("decisions", [])
    if decisions:
        doc.add_heading("Решения и договорённости", level=1)
        for d in decisions:
            bp = doc.add_paragraph(style="List Bullet")
            bp.add_run(d)

    # ---- Сделано (выполненные задачи) --------------------------------------
    done_tasks = analysis.get("done_tasks", [])
    if done_tasks:
        doc.add_heading("Сделано (выполненные задачи)", level=1)
        for task in done_tasks:
            tp = doc.add_paragraph(style="List Bullet")
            tp.add_run(f"☑ {task}")  # ☑ checked box

    # ---- Задачи (нужно сделать) --------------------------------------------
    tasks = analysis.get("tasks", [])
    doc.add_heading("Задачи (нужно сделать)", level=1)
    if tasks:
        for task in tasks:
            tp = doc.add_paragraph(style="List Number")
            tp.add_run(f"☐ {task}")  # ☐ empty checkbox
    else:
        doc.add_paragraph("Задач не выявлено.").paragraph_format.space_after = Pt(6)

    # ---- Мелкие задачи и доработки -----------------------------------------
    minor = analysis.get("minor_tasks", [])
    if minor:
        doc.add_heading("Мелкие задачи и доработки", level=1)
        for task in minor:
            tp = doc.add_paragraph(style="List Bullet")
            tp.add_run(f"☐ {task}")

    # ---- Full transcript on new page ----------------------------------------
    doc.add_page_break()
    doc.add_heading("Полная транскрипция", level=1)

    if segments:
        for seg in segments:
            start = seg.get("start", 0)
            text = seg.get("text", "").strip()
            speaker = seg.get("speaker")
            if not text:
                continue
            line_p = doc.add_paragraph()
            ts_run = line_p.add_run(f"[{_fmt_time(start)}]  ")
            ts_run.font.color.rgb = RGBColor(0x00, 0x70, 0xC0)
            ts_run.font.size = Pt(10)
            if speaker:
                spk_run = line_p.add_run(f"{speaker}: ")
                spk_run.bold = True
            line_p.add_run(text)
            line_p.paragraph_format.space_after = Pt(4)
    else:
        doc.add_paragraph("Транскрипция недоступна.")

    doc.save(str(out_path))
